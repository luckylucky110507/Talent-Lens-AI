import io
import re
from pathlib import Path
from typing import Any


SECTION_ALIASES = {
    "education": ["education", "academic background", "qualifications"],
    "experience": ["experience", "work history", "employment"],
    "projects": ["projects", "personal projects", "academic projects"],
    "certifications": ["certifications", "certificates", "licenses"],
    "skills": ["skills", "technical skills", "core competencies"],
    "summary": ["summary", "profile", "objective", "about me"],
}


def extract_text_from_bytes(filename: str, content: bytes) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if suffix == ".docx":
        from docx import Document
        document = Document(io.BytesIO(content))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(" ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(paragraphs).strip()
    raise ValueError("Unsupported resume format. Please upload a PDF or DOCX file.")


def parse_resume(text: str) -> dict[str, Any]:
    text = (text or "").strip()
    if not text:
        raise ValueError("The resume appears to be empty or could not be read.")
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip()]
    email = _first(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    phone = _first(r"(?:\+?\d[\d\s().-]{8,}\d)", text)
    sections = _sections(lines)
    name = _name(lines, email)
    data = {
        "name": name or "Not detected",
        "email": email or "Not detected",
        "phone": phone or "Not detected",
        "summary": sections.get("summary") or "Not detected",
        "education": sections.get("education") or "Not detected",
        "experience": sections.get("experience") or "Not detected",
        "skills": sections.get("skills") or "Not detected",
        "projects": sections.get("projects") or "Not detected",
        "certifications": sections.get("certifications") or "Not detected",
        "sections_detected": [k for k, v in sections.items() if v],
        "text_length": len(text),
    }
    return data


def _first(pattern: str, text: str) -> str | None:
    match = re.search(pattern, text, re.I)
    return match.group(0).strip() if match else None


def _name(lines: list[str], email: str | None) -> str | None:
    for line in lines[:5]:
        if "@" not in line and not re.search(r"\d", line) and 1 < len(line.split()) <= 5:
            if line.lower() not in {"resume", "curriculum vitae", "cv"}:
                return line
    return None


def _sections(lines: list[str]) -> dict[str, str]:
    markers = {}
    for index, line in enumerate(lines):
        normalized = re.sub(r"[^a-z ]", "", line.lower()).strip()
        for key, aliases in SECTION_ALIASES.items():
            if normalized in aliases or any(normalized.startswith(alias + " ") for alias in aliases):
                markers[index] = key
                break
    result = {}
    ordered = sorted(markers.items())
    for pos, (line_index, key) in enumerate(ordered):
        end = ordered[pos + 1][0] if pos + 1 < len(ordered) else len(lines)
        content = " ".join(lines[line_index + 1:end]).strip()
        if content:
            result[key] = content
    return result
